"""Offline reference-monitor demonstration. Invented local inputs only."""

import argparse
from contextlib import ExitStack
from dataclasses import asdict, dataclass
from datetime import date
from hashlib import sha256
from html import escape
import json
from pathlib import Path
import re
import sys
from xml.etree.ElementTree import Element, ParseError
from zipfile import BadZipFile

from defusedxml import ElementTree as SafeXML
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


NOTICE = "SYNTHETIC DEMONSTRATION - HUMAN REVIEW REQUIRED"
COLUMNS = ("Reference ID", "Name", "Document type", "Revision", "Date")


@dataclass(frozen=True)
class Record:
    source: str
    source_record: str
    reference_id: str
    name: str
    revision: str
    date: str
    context: str

    @property
    def fingerprint(self) -> str:
        return digest(asdict(self))


@dataclass(frozen=True)
class Interest:
    reference_id: str
    name: str
    allow_substring: bool


def digest(value: object) -> str:
    return sha256(json.dumps(value, sort_keys=True, ensure_ascii=True).encode("utf-8")).hexdigest()


def text(value: object, field: str) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > 500:
        raise ValueError(f"{field}: expected nonempty text of at most 500 characters.")
    if any(ord(char) < 32 or 0x7F <= ord(char) <= 0x9F or 0xD800 <= ord(char) <= 0xDFFF
           or char in "\ufffe\uffff" for char in value):
        raise ValueError(f"{field}: invalid control or XML character.")
    return " ".join(value.split())


def identifier(value: object) -> str:
    result = text(value, "identifier")
    if re.fullmatch(r"DEMO-[A-Z0-9]+(?:-[A-Z0-9]+)*", result) is None:
        raise ValueError("Identifiers must use invented uppercase DEMO- values.")
    return result


def iso_date(value: object) -> str:
    result = text(value, "date")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", result) is None:
        raise ValueError("Dates must use YYYY-MM-DD.")
    date.fromisoformat(result)
    return result


def fields(value: object, required: set[str], context: str) -> dict[str, object]:
    if not isinstance(value, dict) or set(value) != required:
        raise ValueError(f"{context}: expected exactly {', '.join(sorted(required))}.")
    return value


def unique_keys(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("Duplicate JSON field.")
        result[key] = value
    return result


def read_json(path: Path) -> object:
    try:
        return json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=unique_keys)
    except (UnicodeError, json.JSONDecodeError) as error:
        raise ValueError("Expected a valid UTF-8 JSON input or checkpoint.") from error


def load_interests(path: Path) -> list[Interest]:
    data = fields(read_json(path), {"synthetic", "interests"}, "Interest list")
    if data["synthetic"] is not True:
        raise ValueError("Interest list must declare synthetic: true.")
    entries = data["interests"]
    if not isinstance(entries, list) or not 1 <= len(entries) <= 100:
        raise ValueError("Supply 1 to 100 invented interests.")
    interests: list[Interest] = []
    seen: set[str] = set()
    for entry in entries:
        item = fields(entry, {"reference_id", "name", "allow_substring"}, "Interest")
        ref = identifier(item["reference_id"])
        if ref in seen or type(item["allow_substring"]) is not bool:
            raise ValueError("Interest IDs must be unique and allow_substring must be Boolean.")
        interests.append(Interest(ref, text(item["name"], "interest name"), item["allow_substring"]))
        seen.add(ref)
    return interests


def policy_key(interests: list[Interest]) -> str:
    return digest([asdict(item) for item in sorted(interests, key=lambda item: item.reference_id)])


def parse_xml(path: Path) -> Element:
    try:
        return SafeXML.fromstring(path.read_bytes(), forbid_dtd=True)
    except (ParseError, DefusedXmlException) as error:
        raise ValueError("Malformed XML/XHTML or forbidden DTD/entity in synthetic source.") from error


def read_edqm(feed: Path) -> list[Record]:
    root = parse_xml(feed)
    if root.tag != "rss" or root.get("synthetic") != "true":
        raise ValueError("Feed must be synthetic RSS.")
    channel = root.find("channel")
    if channel is None:
        raise ValueError("Feed is missing its channel.")
    records: list[Record] = []
    publications: set[str] = set()
    for item in channel.findall("item"):
        publication = identifier(item.findtext("guid"))
        if publication in publications:
            raise ValueError("Duplicate feed publication identifier.")
        publications.add(publication)
        title = text(item.findtext("title"), "publication title")
        link = text(item.findtext("link"), "local announcement filename")
        if re.fullmatch(r"[A-Za-z0-9_-]+\.html", link) is None:
            raise ValueError("Feed links must be local HTML basenames, not URLs or paths.")
        announcement = feed.parent / link
        if announcement.resolve().parent != feed.parent.resolve():
            raise ValueError("Announcement must remain inside the fixture directory.")
        page = parse_xml(announcement)
        if page.tag != "html" or page.get("data-synthetic") != "true":
            raise ValueError("Announcement must be explicitly synthetic XML-compatible HTML.")
        tables = page.findall(".//table[@id='references']")
        if len(tables) != 1:
            raise ValueError("Announcement must contain one references table.")
        rows = tables[0].findall("./tbody/tr")
        if not rows:
            raise ValueError("Announcement references table has no data rows.")
        for row in rows:
            cells = row.findall("td")
            if len(cells) != 4:
                raise ValueError("Each announcement row needs ID, name, revision, and date.")
            values = ["".join(cell.itertext()) for cell in cells]
            ref = identifier(values[0])
            records.append(Record(
                "EDQM-style", f"{publication}/{ref}", ref, text(values[1], "name"),
                text(values[2], "revision"), iso_date(values[3]), title,
            ))
    return records


def read_usp(path: Path) -> list[Record]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ParseError, KeyError, InvalidFileException) as error:
        raise ValueError("Invalid synthetic spreadsheet.") from error
    try:
        if workbook.sheetnames != ["Synthetic changes"]:
            raise ValueError("Workbook must contain only the Synthetic changes sheet.")
        rows = workbook.active.iter_rows()
        header = next(rows, ())
        if tuple(cell.value for cell in header) != COLUMNS:
            raise ValueError("Spreadsheet columns do not match the synthetic schema.")
        records: list[Record] = []
        for row in rows:
            if all(cell.value is None for cell in row):
                continue
            if len(row) != 5 or any(cell.data_type == "f" for cell in row):
                raise ValueError("Spreadsheet rows must have five literal cells, no formulas.")
            ref, name, kind, revision, changed_on = [cell.value for cell in row]
            ref = identifier(ref)
            name = text(name, "name")
            kind = text(kind, "document type")
            revision = text(revision, "revision")
            changed_on = iso_date(changed_on)
            if kind not in {"Monograph", "Chapter"}:
                raise ValueError("Document type must be Monograph or Chapter.")
            if kind == "Monograph":
                records.append(Record("USP-style", ref, ref, name, revision, changed_on,
                                      "Invented change-list spreadsheet"))
        return records
    finally:
        workbook.close()


def match(record: Record, interests: list[Interest]) -> list[dict[str, str]]:
    reasons: list[dict[str, str]] = []
    name = record.name.casefold()
    for interest in interests:
        if record.reference_id == interest.reference_id:
            reason = "reference identifier"
        elif name == interest.name.casefold():
            reason = "exact name"
        elif interest.allow_substring and interest.name.casefold() in name:
            reason = "substring name - verify relevance"
        else:
            continue
        reasons.append({"interest_id": interest.reference_id, "reason": reason})
    return reasons


def load_state(path: Path | None, policy: str) -> set[str]:
    if path is None:
        return set()
    state = fields(read_json(path), {"schema_version", "synthetic", "policy", "observed"}, "Checkpoint")
    if type(state["schema_version"]) is not int or state["schema_version"] != 1 or state["synthetic"] is not True:
        raise ValueError("Unsupported or non-synthetic checkpoint.")
    if state["policy"] != policy:
        raise ValueError("Interest list changed. Start a fresh run without --previous to reassess all records.")
    observed = state["observed"]
    if not isinstance(observed, list) or any(
        not isinstance(key, str) or re.fullmatch(r"[0-9a-f]{64}", key) is None for key in observed
    ):
        raise ValueError("Checkpoint observed values must be SHA-256 fingerprints.")
    if len(set(observed)) != len(observed):
        raise ValueError("Checkpoint contains duplicate fingerprints.")
    return set(observed)


def preview(candidates: list[dict[str, object]], observed: int, new: int) -> str:
    rows = []
    for candidate in candidates:
        record = candidate["record"]
        reasons = candidate["matches"]
        if not isinstance(record, dict) or not isinstance(reasons, list):
            raise ValueError("Invalid review candidate.")
        rows.append("<tr>" + "".join(
            f"<td>{escape(str(record[key]))}</td>"
            for key in ("source", "reference_id", "name", "revision", "date", "context")
        ) + "<td>" + escape("; ".join(str(item["reason"]) for item in reasons)) + "</td></tr>")
    body = "".join(rows) or '<tr><td colspan="7">No new relevant candidates. This does not establish source completeness.</td></tr>'
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>Offline reference monitor</title>
<style>
body {{background:#edf2f7;color:#17283b;font:16px/1.6 system-ui,sans-serif;margin:0}}
main {{max-width:1200px;margin:32px auto;background:white;padding:24px;border-radius:12px}}
.notice {{background:#fff1dc;padding:16px;border-left:5px solid #aa5700}}
.table {{overflow-x:auto}} table {{width:100%;border-collapse:collapse;min-width:800px}}
th,td {{padding:12px;text-align:left;border-bottom:1px solid #ccd7e2;overflow-wrap:anywhere;max-width:260px}}
th {{background:#e7eef6}} a {{color:#164a7b}}
@media(max-width:640px) {{main {{margin:0;padding:16px}}}}
</style></head><body><main>
<p>REGULATORY WORKFLOW DEMO / OFFLINE INPUTS</p><h1>Reference review queue</h1>
<p class="notice">{NOTICE}</p>
<p>{observed} distinct monograph records observed / {new} new records / {len(candidates)} review candidates.</p>
<p>No live access, email delivery, regulatory assessment, or completeness guarantee.</p>
<p><a href="review-queue.json">Structured review queue</a> | <a href="state.json">Local processing checkpoint</a></p>
<div class="table" role="region" aria-label="Review candidates" tabindex="0">
<table><thead><tr><th>Source</th><th>Reference</th><th>Name</th><th>Revision</th><th>Date</th>
<th>Source context</th><th>Match reason</th></tr></thead><tbody>{body}</tbody></table></div>
<p>Substring matches require particular care. State means successful local processing,
not delivered notifications. Use only a completed run's checkpoint for the next run.</p>
</main></body></html>
"""


def run(feed: Path, spreadsheet: Path, interest_path: Path, previous: Path | None, output: Path) -> dict[str, int]:
    interests = load_interests(interest_path)
    policy = policy_key(interests)
    seen = load_state(previous, policy)
    records = read_edqm(feed) + read_usp(spreadsheet)
    unique = {record.fingerprint: record for record in records}
    new = {key: record for key, record in unique.items() if key not in seen}
    candidates: list[dict[str, object]] = []
    for key, record in new.items():
        reasons = match(record, interests)
        if reasons:
            candidates.append({"fingerprint": key, "record": asdict(record), "matches": reasons})
    summary = {"observed": len(unique), "new_records": len(new), "matches": len(candidates),
               "already_seen": len(unique) - len(new), "duplicates": len(records) - len(unique)}
    queue = {"schema_version": 1, "synthetic": True, "summary": summary, "candidates": candidates}
    state = {"schema_version": 1, "synthetic": True, "policy": policy, "observed": sorted(seen | set(unique))}
    contents = {
        "review-queue.json": json.dumps(queue, indent=2, ensure_ascii=True) + "\n",
        "preview.html": preview(candidates, len(unique), len(new)),
        "state.json": json.dumps(state, indent=2) + "\n",
    }
    if output.exists():
        raise ValueError("Output already exists; choose a fresh run directory.")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.mkdir()
    # Keep checkpoints immutable; write this run's state last and roll back failed writes.
    with ExitStack() as rollback:
        rollback.callback(output.rmdir)
        for name, content in contents.items():
            destination = output / name
            with destination.open("x", encoding="utf-8", newline="\n") as stream:
                rollback.callback(destination.unlink)
                stream.write(content)
        rollback.pop_all()
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--feed", type=Path, required=True)
    parser.add_argument("--spreadsheet", type=Path, required=True)
    parser.add_argument("--interests", type=Path, required=True)
    parser.add_argument("--previous", type=Path, help="Checkpoint from a completed previous run.")
    parser.add_argument("--output", type=Path, required=True, help="New run directory.")
    parser.add_argument("--confirm-synthetic", action="store_true", required=True)
    args = parser.parse_args(argv)
    try:
        summary = run(args.feed, args.spreadsheet, args.interests, args.previous, args.output)
    except ValueError as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    except OSError as error:
        print(f"Error: file operation failed ({error.strerror or 'I/O error'}). "
              "Do not use this run; the previous checkpoint was not changed.", file=sys.stderr)
        return 1
    print(f"{NOTICE}. Observed: {summary['observed']}; new: {summary['new_records']}; "
          f"review candidates: {summary['matches']}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
