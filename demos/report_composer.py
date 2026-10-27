"""Compose generic draft documents from explicitly synthetic structured input."""

import argparse
from contextlib import ExitStack
from dataclasses import dataclass
from datetime import date, datetime, timezone
from html import escape
from io import BytesIO
import json
from pathlib import Path
import re
import sys

from docx import Document
from docx.document import Document as WordDocument
from docx.shared import Inches, Pt, RGBColor


NOTICE = "DEMONSTRATION ONLY - NOT FOR REGULATORY SUBMISSION"
REVIEW_STATUS = "DRAFT - NOT APPROVED"
NEUTRAL_AUTHOR = "Regulatory Workflow Demo"
OUTPUT_FILES = ("cover-letter.docx", "annual-report.docx", "preview.html")


@dataclass(frozen=True)
class Amendment:
    sequence: str
    date: date
    description: str


@dataclass(frozen=True)
class Report:
    product: str
    dossier: str
    reporting_year: int
    prepared_on: date
    amendments: tuple[Amendment, ...]


def _fields(value: object, expected: set[str], context: str) -> dict[str, object]:
    if not isinstance(value, dict) or set(value) != expected:
        raise ValueError(f"{context}: supply exactly these fields: {', '.join(sorted(expected))}.")
    return value


def _text(value: object, field: str, limit: int = 200) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ValueError(f"{field}: supply nonempty text of at most {limit} characters.")
    if any(ord(char) < 32 or 0x7F <= ord(char) <= 0x9F
           or 0xD800 <= ord(char) <= 0xDFFF or char in "\ufffe\uffff"
           for char in value):
        raise ValueError(f"{field}: control characters or invalid XML characters are not allowed.")
    return value.strip()


def _identifier(value: object, field: str) -> str:
    text = _text(value, field, 60)
    if re.fullmatch(r"DEMO-[A-Z0-9]+(?:-[A-Z0-9]+)*", text) is None:
        raise ValueError(f"{field}: use an invented DEMO- identifier with uppercase letters or digits.")
    return text


def _date(value: object, field: str) -> date:
    if not isinstance(value, str) or re.fullmatch(r"\d{4}-\d{2}-\d{2}", value) is None:
        raise ValueError(f"{field}: use a valid YYYY-MM-DD date.")
    try:
        return date.fromisoformat(value)
    except ValueError as error:
        raise ValueError(f"{field}: use a valid YYYY-MM-DD date.") from error


def parse_report(value: object) -> Report:
    fields = _fields(
        value,
        {"synthetic", "product", "dossier", "reporting_year", "prepared_on", "amendments"},
        "Report",
    )
    if fields["synthetic"] is not True:
        raise ValueError("synthetic: must be true. Use invented examples only.")
    product = _text(fields["product"], "product")
    dossier = _identifier(fields["dossier"], "dossier")
    year = fields["reporting_year"]
    if type(year) is not int or not 2000 <= year <= 2100:
        raise ValueError("reporting_year: supply an integer between 2000 and 2100.")
    prepared_on = _date(fields["prepared_on"], "prepared_on")
    if prepared_on.year < year:
        raise ValueError("prepared_on: cannot precede the reporting year.")
    records = fields["amendments"]
    if not isinstance(records, list) or len(records) > 100:
        raise ValueError("amendments: supply a list of at most 100 records.")
    amendments: list[Amendment] = []
    seen: set[str] = set()
    for index, record in enumerate(records, start=1):
        context = f"amendments[{index}]"
        row = _fields(record, {"sequence", "date", "description"}, context)
        sequence = _identifier(row["sequence"], f"{context}.sequence")
        if sequence in seen:
            raise ValueError(f"{context}.sequence: duplicate demonstration sequence.")
        changed_on = _date(row["date"], f"{context}.date")
        if changed_on.year != year or changed_on > prepared_on:
            raise ValueError(f"{context}.date: must be in the reporting year and on or before prepared_on.")
        amendments.append(Amendment(sequence, changed_on, _text(row["description"], f"{context}.description", 1000)))
        seen.add(sequence)
    return Report(product, dossier, year, prepared_on, tuple(amendments))


def _unique_keys(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("Input JSON contains a duplicate field.")
        result[key] = value
    return result


def load_report(path: Path) -> Report:
    try:
        value: object = json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=_unique_keys)
    except json.JSONDecodeError as error:
        raise ValueError(f"Invalid JSON at line {error.lineno}, column {error.colno}.") from error
    except UnicodeError as error:
        raise ValueError("Input must be a UTF-8 JSON file.") from error
    return parse_report(value)


def _document(title: str, report: Report) -> WordDocument:
    document = Document()
    properties = document.core_properties
    properties.author = NEUTRAL_AUTHOR
    properties.last_modified_by = NEUTRAL_AUTHOR
    properties.title = title
    properties.subject = NOTICE
    properties.comments = "Generated from invented data using a generic demonstration layout."
    properties.keywords = "synthetic, demonstration, draft"
    properties.identifier = report.dossier
    properties.revision = 1
    timestamp = datetime.combine(report.prepared_on, datetime.min.time(), tzinfo=timezone.utc)
    properties.created = timestamp
    properties.modified = timestamp
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.paragraph_format.space_after = Pt(8)
    section.header.paragraphs[0].text = NOTICE
    section.footer.paragraphs[0].text = REVIEW_STATUS
    document.add_heading(title, 0)
    notice = document.add_paragraph()
    run = notice.add_run(NOTICE)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("9C3328")
    document.add_paragraph(REVIEW_STATUS)
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in (
        ("Product", report.product),
        ("Dossier", report.dossier),
        ("Reporting year", str(report.reporting_year)),
        ("Prepared on", report.prepared_on.isoformat()),
    ):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    return document


def _cover_letter(report: Report) -> WordDocument:
    document = _document("Draft cover letter", report)
    document.add_heading("Illustrative purpose", 1)
    document.add_paragraph(
        "This fictional cover letter accompanies a demonstration annual report. "
        "It is not correspondence to an authority and is not intended for submission."
    )
    document.add_paragraph(
        f"The companion draft contains {len(report.amendments)} amendment record(s) "
        "supplied in the synthetic input. No completeness or compliance assessment has been performed."
    )
    document.add_heading("Review required", 1)
    document.add_paragraph(
        "Check all input values and the companion report. This demonstration does not "
        "include a signature, approval, or submission step."
    )
    return document


def _annual_report(report: Report) -> WordDocument:
    document = _document("Draft annual report", report)
    document.add_heading("Scope", 1)
    document.add_paragraph(
        "This generic layout illustrates structured document assembly, not an official "
        "regulatory template. All records must be invented."
    )
    document.add_heading("Amendments supplied for review", 1)
    if report.amendments:
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Shading Accent 1"
        for cell, label in zip(table.rows[0].cells, ("Sequence", "Date", "Description")):
            cell.text = label
        for amendment in report.amendments:
            cells = table.add_row().cells
            for cell, text in zip(cells, (amendment.sequence, amendment.date.isoformat(), amendment.description)):
                cell.text = text
    else:
        document.add_paragraph("No amendments supplied in the synthetic input. Completeness has not been assessed.")
    document.add_heading("Reviewer checklist", 1)
    for text in (
        "Verify product, dossier, reporting year, and preparation date.",
        "Compare each amendment row with the supplied input.",
        "Do not interpret this draft as approved, complete, or submission-ready.",
    ):
        document.add_paragraph(text, style="List Bullet")
    return document


def render_preview(report: Report) -> str:
    rows = "".join(
        f"<tr><td>{escape(item.sequence)}</td><td>{item.date.isoformat()}</td>"
        f"<td>{escape(item.description)}</td></tr>"
        for item in report.amendments
    )
    if not rows:
        rows = '<tr><td colspan="3">No amendments supplied. Completeness has not been assessed.</td></tr>'
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="author" content="{NEUTRAL_AUTHOR}">
<title>Synthetic annual report preview</title>
<style>
body {{ margin: 0; background: #edf2f7; color: #17283b; font: 16px/1.6 system-ui, sans-serif; }}
main {{ max-width: 960px; margin: 40px auto; padding: 32px; background: white; border-radius: 12px; }}
h1, h2 {{ line-height: 1.2; }}
.notice {{ padding: 16px; background: #fff1dc; border-left: 5px solid #aa5700; }}
.status {{ color: #9c3328; font-weight: bold; }}
dl {{ display: grid; grid-template-columns: 150px 1fr; gap: 8px 20px; }}
dt {{ font-weight: bold; }} dd {{ margin: 0; overflow-wrap: anywhere; }}
table {{ width: 100%; border-collapse: collapse; }}
th, td {{ text-align: left; padding: 12px; border-bottom: 1px solid #ccd7e2; overflow-wrap: anywhere; }}
th {{ background: #e7eef6; }}
nav a {{ display: inline-block; margin: 4px 16px 4px 0; color: #164a7b; }}
@media (max-width: 640px) {{ main {{ margin: 0; padding: 16px; }} dl {{ grid-template-columns: 1fr; }} }}
@media print {{ body {{ background: white; }} main {{ margin: 0; }} }}
</style>
</head>
<body><main>
<p>REGULATORY WORKFLOW DEMO / SYNTHETIC DATA</p>
<h1>Annual report composer</h1>
<p class="notice">{NOTICE}</p>
<p class="status">{REVIEW_STATUS}</p>
<dl>
<dt>Product</dt><dd>{escape(report.product)}</dd>
<dt>Dossier</dt><dd>{escape(report.dossier)}</dd>
<dt>Reporting year</dt><dd>{report.reporting_year}</dd>
<dt>Prepared on</dt><dd>{report.prepared_on.isoformat()}</dd>
</dl>
<nav aria-label="Generated documents">
<a href="cover-letter.docx">Download draft cover letter</a>
<a href="annual-report.docx">Download draft annual report</a>
</nav>
<h2>Cover letter overview</h2>
<p>Fictional companion letter for a draft containing {len(report.amendments)} amendment record(s).
Not correspondence to an authority. No signature or approval is provided.</p>
<h2>Amendments supplied for review</h2>
<table><thead><tr><th scope="col">Sequence</th><th scope="col">Date</th>
<th scope="col">Description</th></tr></thead><tbody>{rows}</tbody></table>
<h2>Before any use</h2>
<p>Check all values against the supplied synthetic input. This preview and the
companion documents are not official templates, regulatory advice, or evidence
of completeness or compliance. Nothing is submitted or approved.</p>
</main></body></html>
"""


def generate(report: Report, output: Path) -> None:
    if output.exists():
        raise ValueError("Output already exists. Choose a new output directory; nothing was overwritten.")
    cover = BytesIO()
    annual = BytesIO()
    _cover_letter(report).save(cover)
    _annual_report(report).save(annual)
    contents = (cover.getvalue(), annual.getvalue(), render_preview(report).encode("utf-8"))
    output.parent.mkdir(parents=True, exist_ok=True)
    # Reserve the destination exclusively and roll back only files created here.
    output.mkdir()
    with ExitStack() as rollback:
        rollback.callback(output.rmdir)
        for name, content in zip(OUTPUT_FILES, contents):
            path = output / name
            with path.open("xb") as stream:
                rollback.callback(path.unlink)
                stream.write(content)
        rollback.pop_all()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="Invented JSON input only.")
    parser.add_argument("--output", type=Path, required=True, help="New directory for all three outputs.")
    parser.add_argument(
        "--confirm-synthetic", action="store_true", required=True,
        help="Confirm that the input contains invented data, not operational records.",
    )
    args = parser.parse_args(argv)
    try:
        report = load_report(args.input)
        generate(report, args.output)
    except ValueError as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    except OSError as error:
        print(f"Error: file operation failed ({error.strerror or 'I/O error'}). No completed package was published.",
              file=sys.stderr)
        return 1
    print(f"{REVIEW_STATUS}. Created {', '.join(OUTPUT_FILES)} in the requested output directory.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
