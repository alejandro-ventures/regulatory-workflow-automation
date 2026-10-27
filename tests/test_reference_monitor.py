from contextlib import redirect_stderr, redirect_stdout
from dataclasses import replace
from io import BytesIO, StringIO
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

from demos.monitor_fixtures import generate, spreadsheet
from demos.reference_monitor import (
    Interest, NOTICE, load_interests, load_state, main, match, policy_key,
    read_edqm, read_usp, run,
)


FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "monitoring"


class ReferenceMonitorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "fixtures"
        shutil.copytree(FIXTURES, self.source)
        self.feed = self.source / "feed.xml"
        self.interests = self.source / "interests.json"
        self.exports = self.root / "exports"
        generate(self.exports)
        self.baseline = self.exports / "baseline.xlsx"
        self.revised = self.exports / "revised.xlsx"

    def execute(self, name: str, previous: Path | None = None, revised: bool = False) -> dict[str, int]:
        return run(self.feed, self.revised if revised else self.baseline,
                   self.interests, previous, self.root / name)

    def test_first_repeat_revised_runs_and_output_shape(self) -> None:
        with patch("socket.socket", side_effect=AssertionError("Network not allowed")):
            first = self.execute("first")
            checkpoint = self.root / "first" / "state.json"
            original = checkpoint.read_bytes()
            second = self.execute("second", checkpoint)
            third = self.execute("third", self.root / "second" / "state.json", True)
        self.assertEqual(first, {"observed": 5, "new_records": 5, "matches": 3,
                                 "already_seen": 0, "duplicates": 0})
        self.assertEqual(second, {"observed": 5, "new_records": 0, "matches": 0,
                                  "already_seen": 5, "duplicates": 0})
        self.assertEqual(third, {"observed": 5, "new_records": 1, "matches": 1,
                                 "already_seen": 4, "duplicates": 0})
        self.assertEqual(checkpoint.read_bytes(), original)
        queue = json.loads((self.root / "first" / "review-queue.json").read_text())
        self.assertEqual(set(queue), {"schema_version", "synthetic", "summary", "candidates"})
        self.assertIs(queue["synthetic"], True)
        self.assertEqual([item["matches"][0]["reason"] for item in queue["candidates"]],
                         ["reference identifier", "exact name", "substring name - verify relevance"])
        for item in queue["candidates"]:
            self.assertEqual(set(item), {"fingerprint", "record", "matches"})
            self.assertEqual(set(item["record"]), {
                "source", "source_record", "reference_id", "name", "revision", "date", "context",
            })
            self.assertEqual(len(item["fingerprint"]), 64)
        self.assertEqual(len(json.loads((self.root / "third" / "state.json").read_text())["observed"]), 6)
        self.assertEqual({path.name for path in (self.root / "first").iterdir()},
                         {"preview.html", "review-queue.json", "state.json"})
        html = (self.root / "first" / "preview.html").read_text()
        self.assertIn(NOTICE, html)
        self.assertEqual(html.count("<tr>"), 4)
        self.assertNotIn(str(self.root), html)

    def test_source_filter_and_matching(self) -> None:
        records = read_edqm(self.feed) + read_usp(self.baseline)
        self.assertEqual(len(records), 5)
        self.assertNotIn("DEMO-REF-F", [record.reference_id for record in records])
        record = replace(records[0], reference_id="DEMO-OTHER", name="example substance alpha")
        interest = Interest("DEMO-REF-A", "Example Substance Alpha", False)
        self.assertEqual(match(record, [interest])[0]["reason"], "exact name")
        self.assertEqual(match(replace(record, name="example substance alpha hydrate"), [interest]), [])
        self.assertEqual(len(match(replace(record, name="example substance alpha hydrate"),
                                   [replace(interest, allow_substring=True)])), 1)

    def test_duplicates_collapsed_and_unmatched_record_observed(self) -> None:
        original = read_edqm(self.feed)
        with patch("demos.reference_monitor.read_edqm", return_value=original + [original[0]]):
            summary = self.execute("duplicates")
        self.assertEqual(summary["duplicates"], 1)
        self.assertEqual(summary["matches"], 3)
        state = json.loads((self.root / "duplicates" / "state.json").read_text())
        self.assertIn(original[1].fingerprint, state["observed"])

    def test_same_id_with_changed_content_is_new(self) -> None:
        self.execute("first")
        original = read_edqm(self.feed)
        changed = [replace(original[0], revision="Revision 2"), original[1]]
        with patch("demos.reference_monitor.read_edqm", return_value=changed):
            result = self.execute("second", self.root / "first" / "state.json")
        self.assertEqual(result["new_records"], 1)
        self.assertEqual(result["matches"], 1)

    def test_corrupt_missing_and_wrong_version_checkpoints_fail(self) -> None:
        path = self.root / "broken.json"
        policy = policy_key(load_interests(self.interests))
        invalid = [
            "{",
            json.dumps({"schema_version": 1}),
            json.dumps({"schema_version": True, "synthetic": True, "policy": policy, "observed": []}),
            json.dumps({"schema_version": 2, "synthetic": True, "policy": policy, "observed": []}),
            json.dumps({"schema_version": 1, "synthetic": False, "policy": policy, "observed": []}),
            json.dumps({"schema_version": 1, "synthetic": True, "policy": policy, "observed": ["bad"]}),
            json.dumps({"schema_version": 1, "synthetic": True, "policy": policy, "observed": ["a" * 64] * 2}),
        ]
        for value in invalid:
            path.write_text(value, encoding="utf-8")
            with self.subTest(value=value), self.assertRaises(ValueError):
                load_state(path, policy)
        with self.assertRaises(FileNotFoundError):
            load_state(self.root / "missing.json", policy)

    def test_changed_interest_policy_requires_explicit_fresh_run(self) -> None:
        self.execute("first")
        data = json.loads(self.interests.read_text())
        data["interests"][1]["allow_substring"] = False
        self.interests.write_text(json.dumps(data))
        with self.assertRaisesRegex(ValueError, "Interest list changed"):
            self.execute("second", self.root / "first" / "state.json")
        self.assertFalse((self.root / "second").exists())
        self.assertEqual(self.execute("fresh")["matches"], 2)

    def test_invalid_interest_lists(self) -> None:
        data = json.loads(self.interests.read_text())
        cases = [
            {"synthetic": False, "interests": data["interests"]},
            {"synthetic": True, "interests": []},
            {"synthetic": True, "interests": [data["interests"][0]] * 2},
            {"synthetic": True, "interests": [{"reference_id": "DEMO-A", "name": "", "allow_substring": False}]},
            {"synthetic": True, "interests": [{"reference_id": "DEMO-A", "name": "Example", "allow_substring": "yes"}]},
        ]
        for value in cases:
            self.interests.write_text(json.dumps(value))
            with self.subTest(value=value), self.assertRaises(ValueError):
                load_interests(self.interests)
        self.interests.write_text('{"synthetic": true, "synthetic": true, "interests": []}')
        with self.assertRaisesRegex(ValueError, "Duplicate"):
            load_interests(self.interests)

    def test_feed_rejects_malformed_xml_entities_remote_and_traversal_links(self) -> None:
        original = self.feed.read_text()
        cases = [
            "<rss>",
            '<!DOCTYPE rss [<!ENTITY x "test">]><rss synthetic="true"><channel/></rss>',
            original.replace('synthetic="true"', 'synthetic="false"'),
            original.replace("announcement.html", "https://example.invalid/announcement.html"),
            original.replace("announcement.html", "../announcement.html"),
            original.replace("announcement.html", "..\\announcement.html"),
        ]
        for content in cases:
            self.feed.write_text(content)
            with self.subTest(content=content), self.assertRaises(ValueError):
                read_edqm(self.feed)

    def test_missing_and_broken_announcement_are_not_silently_skipped(self) -> None:
        path = self.source / "announcement.html"
        original = path.read_text()
        for content in (
            original.replace('id="references"', 'id="other"'),
            original.replace('data-synthetic="true"', 'data-synthetic="false"'),
            original.replace("<td>Revision 1</td>", ""),
            original.replace("2030-01-15", "2030-02-30"),
        ):
            path.write_text(content)
            with self.subTest(content=content), self.assertRaises(ValueError):
                self.execute("invalid")
            self.assertFalse((self.root / "invalid").exists())
        path.unlink()
        with self.assertRaises(FileNotFoundError):
            read_edqm(self.feed)

    def test_spreadsheet_schema_formula_date_and_type_errors(self) -> None:
        for cell, value in (("A1", "Unexpected"), ("B2", "=1+1"), ("E2", "2030-02-30"),
                            ("C2", "Unknown"), ("A2", "1234"), ("B2", None)):
            workbook = load_workbook(BytesIO(spreadsheet()))
            workbook.active[cell] = value
            workbook.save(self.baseline)
            workbook.close()
            with self.subTest(cell=cell, value=value), self.assertRaises(ValueError):
                read_usp(self.baseline)
        self.baseline.write_bytes(b"not a workbook")
        with self.assertRaisesRegex(ValueError, "Invalid"):
            read_usp(self.baseline)

    def test_empty_valid_sources(self) -> None:
        self.feed.write_text('<rss synthetic="true"><channel/></rss>')
        workbook = load_workbook(BytesIO(spreadsheet()))
        workbook.active.delete_rows(2, 4)
        workbook.save(self.baseline)
        workbook.close()
        self.assertEqual(self.execute("empty")["observed"], 0)
        self.assertIn("No new relevant candidates", (self.root / "empty" / "preview.html").read_text())

    def test_output_collision_preserves_all_files(self) -> None:
        self.execute("first")
        before = {path.name: path.read_bytes() for path in (self.root / "first").iterdir()}
        with self.assertRaisesRegex(ValueError, "already exists"):
            self.execute("first")
        self.assertEqual(before, {path.name: path.read_bytes() for path in (self.root / "first").iterdir()})

    def test_failures_preserve_previous_checkpoint_and_remove_partial_output(self) -> None:
        self.execute("first")
        checkpoint = self.root / "first" / "state.json"
        original = checkpoint.read_bytes()
        original_open = Path.open
        for failed_name in ("review-queue.json", "preview.html", "state.json"):
            def failing_open(path, *args, **kwargs):
                if path.parent.name == "failed" and path.name == failed_name:
                    raise OSError("Simulated output failure")
                return original_open(path, *args, **kwargs)

            with self.subTest(failed_name=failed_name):
                with patch("pathlib.Path.open", autospec=True, side_effect=failing_open):
                    with self.assertRaisesRegex(OSError, "Simulated"):
                        self.execute("failed", checkpoint, True)
                self.assertEqual(checkpoint.read_bytes(), original)
                self.assertFalse((self.root / "failed").exists())
        self.assertEqual(self.execute("retry", checkpoint, True)["matches"], 1)

    def test_render_error_writes_nothing(self) -> None:
        with patch("demos.reference_monitor.preview", side_effect=ValueError("Invalid render")):
            with self.assertRaisesRegex(ValueError, "Invalid render"):
                self.execute("bad")
        self.assertFalse((self.root / "bad").exists())

    def test_html_escapes_all_source_text(self) -> None:
        records = read_edqm(self.feed)
        records[0] = replace(records[0], name="<script>example</script>", context="Demo <b>context</b>")
        with patch("demos.reference_monitor.read_edqm", return_value=records):
            self.execute("escaped")
        html = (self.root / "escaped" / "preview.html").read_text()
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertIn("Demo &lt;b&gt;context&lt;/b&gt;", html)

    def test_cli_success_error_and_confirmation(self) -> None:
        args = ["--feed", str(self.feed), "--spreadsheet", str(self.baseline),
                "--interests", str(self.interests), "--output", str(self.root / "cli"),
                "--confirm-synthetic"]
        stdout, stderr = StringIO(), StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            self.assertEqual(main(args), 0)
        self.assertIn("review candidates: 3", stdout.getvalue())
        self.assertEqual(stderr.getvalue(), "")
        stdout, stderr = StringIO(), StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            self.assertEqual(main(args), 1)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("already exists", stderr.getvalue())
        with redirect_stderr(StringIO()), self.assertRaises(SystemExit) as caught:
            main(args[:-1])
        self.assertEqual(caught.exception.code, 2)
        self.feed.unlink()
        stdout, stderr = StringIO(), StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            self.assertEqual(main(args), 1)
        self.assertIn("file operation failed", stderr.getvalue())
        self.assertNotIn(str(self.root), stderr.getvalue())

    def test_fixture_metadata_and_overwrite_protection(self) -> None:
        workbook = load_workbook(self.baseline)
        self.assertEqual(workbook.properties.creator, "Regulatory Workflow Demo")
        self.assertEqual(workbook.properties.lastModifiedBy, "Regulatory Workflow Demo")
        workbook.close()
        with ZipFile(self.baseline) as archive:
            self.assertFalse(any("externalLink" in name for name in archive.namelist()))
        with self.assertRaisesRegex(ValueError, "already exists"):
            generate(self.exports)


if __name__ == "__main__":
    unittest.main()
