from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
from datetime import date
from io import StringIO
import json
from pathlib import Path
import tempfile
from typing import IO
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document

from demos.report_composer import (
    NEUTRAL_AUTHOR,
    NOTICE,
    OUTPUT_FILES,
    REVIEW_STATUS,
    generate,
    load_report,
    main,
    parse_report,
    render_preview,
)


FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "reporting" / "annual-report.json"


class ReportComposerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.data = json.loads(FIXTURE.read_text(encoding="utf-8"))
        self.report = load_report(FIXTURE)

    def test_fixture_parses(self) -> None:
        self.assertEqual(self.report.product, "Example Material Alpha")
        self.assertEqual(self.report.prepared_on, date(2031, 1, 15))
        self.assertEqual(len(self.report.amendments), 2)

    def test_invalid_inputs(self) -> None:
        cases = [
            {"synthetic": False},
            {"synthetic": 1},
            {"product": ""},
            {"product": "Bad\x00text"},
            {"product": "Bad\x7ftext"},
            {"product": "Bad\x85text"},
            {"product": "Bad\ud800text"},
            {"product": "Bad\ufffftext"},
            {"product": 12},
            {"product": "x" * 201},
            {"dossier": "123456"},
            {"reporting_year": True},
            {"reporting_year": "2030"},
            {"reporting_year": 1999},
            {"prepared_on": "2031-02-30"},
            {"prepared_on": "20310115"},
            {"prepared_on": "2029-12-31"},
            {"amendments": {}},
            {"amendments": [{}]},
            {"amendments": [None]},
            {"extra": "not allowed"},
        ]
        for changes in cases:
            with self.subTest(changes=changes):
                with self.assertRaises(ValueError):
                    parse_report(self.data | changes)
        for invalid in (None, [], "text", 42):
            with self.subTest(root=invalid):
                with self.assertRaises(ValueError):
                    parse_report(invalid)
        missing = dict(self.data)
        del missing["product"]
        with self.assertRaises(ValueError):
            parse_report(missing)

    def test_invalid_amendments(self) -> None:
        for changes in (
            {"date": "2031-01-01"},
            {"sequence": "0001"},
            {"description": ""},
            {"date": "2030-02-30"},
        ):
            data = deepcopy(self.data)
            data["amendments"][0].update(changes)
            with self.subTest(changes=changes):
                with self.assertRaises(ValueError):
                    parse_report(data)
        duplicate = self.data | {"amendments": [self.data["amendments"][0]] * 2}
        with self.assertRaisesRegex(ValueError, "duplicate"):
            parse_report(duplicate)
        with self.assertRaisesRegex(ValueError, "at most 100"):
            parse_report(self.data | {"amendments": self.data["amendments"] * 51})
        future = deepcopy(self.data)
        future["prepared_on"] = "2030-01-01"
        with self.assertRaisesRegex(ValueError, "on or before"):
            parse_report(future)

    def test_invalid_json_and_encoding(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "input.json"
            for content, message in (
                (b'{"broken":', "Invalid JSON"),
                (b'{"synthetic": true, "synthetic": false}', "duplicate"),
                (b"\xff", "UTF-8"),
            ):
                path.write_bytes(content)
                with self.subTest(content=content):
                    with self.assertRaisesRegex(ValueError, message):
                        load_report(path)

    def test_package_content_tables_and_metadata(self) -> None:
        before = FIXTURE.read_bytes()
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            generate(self.report, output)
            self.assertEqual({path.name for path in output.iterdir()}, set(OUTPUT_FILES))
            for name in OUTPUT_FILES[:2]:
                document = Document(str(output / name))
                paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
                self.assertIn(NOTICE, paragraphs)
                self.assertIn(REVIEW_STATUS, paragraphs)
                self.assertIn(NOTICE, document.sections[0].header.paragraphs[0].text)
                self.assertIn(REVIEW_STATUS, document.sections[0].footer.paragraphs[0].text)
                self.assertEqual(document.core_properties.author, NEUTRAL_AUTHOR)
                self.assertEqual(document.core_properties.last_modified_by, NEUTRAL_AUTHOR)
                self.assertEqual(document.core_properties.created.date(), self.report.prepared_on)
                self.assertEqual(document.core_properties.modified.date(), self.report.prepared_on)
                self.assertEqual(document.tables[0].cell(0, 1).text, self.report.product)
                with ZipFile(output / name) as archive:
                    self.assertFalse(any("externalLink" in item for item in archive.namelist()))
            annual = Document(str(output / "annual-report.docx"))
            table = annual.tables[1]
            self.assertEqual(len(table.rows), 1 + len(self.report.amendments))
            for row, amendment in zip(table.rows[1:], self.report.amendments):
                self.assertEqual(
                    [cell.text for cell in row.cells],
                    [amendment.sequence, amendment.date.isoformat(), amendment.description],
                )
            preview = (output / "preview.html").read_text(encoding="utf-8")
            for text in (NOTICE, REVIEW_STATUS, self.report.product, 'href="annual-report.docx"'):
                self.assertIn(text, preview)
            self.assertEqual(preview.count("<tr>"), 3)
        self.assertEqual(FIXTURE.read_bytes(), before)

    def test_html_escapes_input_and_docx_preserves_text(self) -> None:
        product = '<script>alert("invented")</script> & Example'
        report = parse_report(self.data | {"product": product})
        preview = render_preview(report)
        self.assertNotIn("<script>", preview)
        self.assertIn("&lt;script&gt;", preview)
        self.assertIn("&amp; Example", preview)
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            generate(report, output)
            document = Document(str(output / "annual-report.docx"))
            self.assertEqual(document.tables[0].cell(0, 1).text, product)

    def test_no_amendments(self) -> None:
        report = parse_report(self.data | {"amendments": []})
        self.assertIn("No amendments supplied", render_preview(report))
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            generate(report, output)
            annual = Document(str(output / "annual-report.docx"))
            self.assertEqual(len(annual.tables), 1)
            self.assertTrue(any("No amendments supplied" in paragraph.text for paragraph in annual.paragraphs))

    def test_existing_output_is_not_overwritten(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            generate(self.report, output)
            original = {path.name: path.read_bytes() for path in output.iterdir()}
            with self.assertRaisesRegex(ValueError, "already exists"):
                generate(self.report, output)
            self.assertEqual({path.name: path.read_bytes() for path in output.iterdir()}, original)

    def test_render_failure_does_not_publish_partial_package(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            with patch("demos.report_composer.render_preview", side_effect=OSError("test failure")):
                with self.assertRaises(OSError):
                    generate(self.report, output)
            self.assertFalse(output.exists())
            self.assertEqual(list(Path(temporary).iterdir()), [])

    def test_write_failure_removes_own_partial_files(self) -> None:
        original_open = Path.open

        def fail_second_file(path: Path, mode: str) -> IO[bytes]:
            if path.name == "annual-report.docx":
                raise OSError("test write failure")
            return original_open(path, mode)

        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            with patch("demos.report_composer.Path.open", autospec=True, side_effect=fail_second_file):
                with self.assertRaisesRegex(OSError, "test write failure"):
                    generate(self.report, output)
            self.assertFalse(output.exists())
            self.assertEqual(list(Path(temporary).iterdir()), [])

    def test_cli_success_failure_and_confirmation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            arguments = ["--input", str(FIXTURE), "--output", str(output), "--confirm-synthetic"]
            stdout, stderr = StringIO(), StringIO()
            with redirect_stdout(stdout), redirect_stderr(stderr):
                self.assertEqual(main(arguments), 0)
            self.assertIn("Created", stdout.getvalue())
            self.assertEqual(stderr.getvalue(), "")
            stdout, stderr = StringIO(), StringIO()
            with redirect_stdout(stdout), redirect_stderr(stderr):
                self.assertEqual(main(arguments), 1)
            self.assertEqual(stdout.getvalue(), "")
            self.assertIn("already exists", stderr.getvalue())
            with redirect_stderr(StringIO()), self.assertRaises(SystemExit) as caught:
                main(arguments[:-1])
            self.assertEqual(caught.exception.code, 2)

    def test_cli_missing_input_is_explicit_and_does_not_leak_path(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            stdout, stderr = StringIO(), StringIO()
            root = Path(temporary)
            with redirect_stdout(stdout), redirect_stderr(stderr):
                result = main([
                    "--input", str(root / "missing.json"), "--output", str(root / "out"),
                    "--confirm-synthetic",
                ])
            self.assertEqual(result, 1)
            self.assertEqual(stdout.getvalue(), "")
            self.assertIn("file operation failed", stderr.getvalue())
            self.assertNotIn(temporary, stderr.getvalue())
            self.assertFalse((root / "out").exists())

    def test_cli_invalid_input_does_not_create_output(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "invalid.json"
            path.write_text(json.dumps(self.data | {"dossier": "123456"}), encoding="utf-8")
            stdout, stderr = StringIO(), StringIO()
            with redirect_stdout(stdout), redirect_stderr(stderr):
                result = main([
                    "--input", str(path), "--output", str(root / "out"), "--confirm-synthetic",
                ])
            self.assertEqual(result, 1)
            self.assertEqual(stdout.getvalue(), "")
            self.assertIn("dossier:", stderr.getvalue())
            self.assertFalse((root / "out").exists())

    def test_cli_write_failure_has_no_success_message(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            stdout, stderr = StringIO(), StringIO()
            with patch("demos.report_composer.render_preview", side_effect=OSError("test failure")):
                with redirect_stdout(stdout), redirect_stderr(stderr):
                    result = main([
                        "--input", str(FIXTURE), "--output", str(root / "out"), "--confirm-synthetic",
                    ])
            self.assertEqual(result, 1)
            self.assertEqual(stdout.getvalue(), "")
            self.assertIn("No completed package was published", stderr.getvalue())
            self.assertEqual(list(root.iterdir()), [])


if __name__ == "__main__":
    unittest.main()
